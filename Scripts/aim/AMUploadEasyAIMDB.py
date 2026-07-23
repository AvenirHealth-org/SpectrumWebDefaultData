import os
import math
import ujson
from openpyxl import load_workbook

from AvenirCommon.Database import GB_upload_json, GB_get_db_json, GB_upload_file
from AvenirCommon.Logger import log
from AvenirCommon.Util import formatCountryFName, GBRange
from AvenirCommon.Logger import log

from DefaultDataUtil import *

from SpectrumCommon.Const.GB import GB_Nan
from SpectrumCommon.Modvars.GB.GBUtil import get_country_ISO3Alpha, get_country_details


easyaim_json_path = os.path.join(os.getcwd(), 'JSONData', 'aim', 'easyAIM')


def _is_empty(value):
    return value is None or (isinstance(value, float) and math.isnan(value))


def _get_end_col_from_header(header_row, start_col_index):
    for col_index in range(len(header_row) - 1, start_col_index - 1, -1):
        if not _is_empty(header_row[col_index]):
            return col_index
    return start_col_index


def isValueByYearSheet(sheetName):
    return not (
        sheetName in [
            'PercentHIVPopEligible',
            'PercNotBFNotRecARVs',
            'PercNotBFRecARVs',
            'LocalAdjustmentFactor',
            'NewARTPatAllocation',
        ]
    )


def getValuesByYear(countries, sheet, sheetName, startRow=3):
    startCol = 4
    header_row = next(sheet.iter_rows(min_row=3, max_row=3, values_only=True), ())
    endCol = _get_end_col_from_header(header_row, startCol)
    startYear = int(header_row[startCol])
    endYear = int(header_row[endCol])

    for row_values in sheet.iter_rows(
        min_row=startRow + 1,
        max_row=sheet.max_row,
        min_col=1,
        max_col=endCol + 1,
        values_only=True,
    ):
        countryCodeValue = row_values[0] if len(row_values) > 0 else None
        countryName = row_values[1] if len(row_values) > 1 else None

        if (not _is_empty(countryCodeValue)) and (not _is_empty(countryName)):
            countryCode = int(countryCodeValue)
            values = [
                0.0 if _is_empty(value) else float(value)
                for value in row_values[startCol : endCol + 1]
            ]

            data = {
                'countryName': countryName,
                'countryCode': countryCode,
                'startYear': startYear,
                'endYear': endYear,
                'values': values,
            }

            addDataByCountryCode(countryCode, countries, sheetName, data)


def getValuesNotByYear(countries, sheet, sheetName):
    startCol = 4
    header_row = next(sheet.iter_rows(min_row=3, max_row=3, values_only=True), ())
    endCol = _get_end_col_from_header(header_row, startCol)

    for row_values in sheet.iter_rows(
        min_row=4,
        max_row=sheet.max_row,
        min_col=1,
        max_col=endCol + 1,
        values_only=True,
    ):
        countryCodeValue = row_values[0] if len(row_values) > 0 else None
        countryName = row_values[1] if len(row_values) > 1 else None

        if (not _is_empty(countryCodeValue)) and (not _is_empty(countryName)):
            countryCode = int(countryCodeValue)
            data = {
                'countryName': countryName,
                'countryCode': countryCode,
            }

            for col in GBRange(startCol, endCol):
                headerValue = header_row[col] if len(header_row) > col else None
                rowValue = row_values[col] if len(row_values) > col else None
                dataName = 'value' if _is_empty(headerValue) else headerValue
                data[dataName] = 0 if _is_empty(rowValue) else float(rowValue)

            addDataByCountryCode(countryCode, countries, sheetName, data)


def write_easyAIM_db(version, country=''):
    FQName = os.path.join(os.getcwd(), 'SourceData', 'aim', 'EasyAIMData.xlsx')

    countries = {}
    xlsx = load_workbook(FQName, data_only=True, read_only=True)
    for sheetName in xlsx.sheetnames:
        sheet = xlsx[sheetName]

        if isValueByYearSheet(sheetName):
            getValuesByYear(countries, sheet, sheetName)
        else:
            getValuesNotByYear(countries, sheet, sheetName)

    xlsx.close()

    os.makedirs(easyaim_json_path, exist_ok=True)
    for countryCode in countries:
        countryData = countries[countryCode]
        ISO3_Alpha = get_country_ISO3Alpha(countryData['countryCode'])

        if ISO3_Alpha != 'notFound':
            log('Writing ' + get_country_details(ISO3_Alpha)['name'])
            FName = formatCountryFName(ISO3_Alpha, version)
            with open(os.path.join(easyaim_json_path, FName), 'w') as f:
                ujson.dump(countryData, f)
            # Sample country copies Kenya
            if ISO3_Alpha == 'KEN':
                log('Writing Sample Country')
                FName = formatCountryFName('SAMP', version)
                with open(os.path.join(easyaim_json_path, FName), 'w') as f:
                    ujson.dump(countryData, f)


def upload_easyAIM_db(version):
    uploadFilesInDir('aim', easyaim_json_path, version, pathMod='easyAIM/')
